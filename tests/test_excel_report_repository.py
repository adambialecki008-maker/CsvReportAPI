from excel_report_repository import create_excel_report
from csv_file_repository import read_csv_file
from summary import build_summary_data
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

sample_file = "tests/fixtures/sample_sales.csv"
data = read_csv_file(sample_file)
summary = build_summary_data(data)


def test_create_excel_report_contains_expected_sheets():
    report = create_excel_report(data, summary)
    workbook = load_workbook(report)
    assert workbook.sheetnames == ["Sales", "Summary"]
    assert workbook["Sales"].max_row == 13  # header + 12 rekordów


def test_create_excel_report_contains_expected_headers_in_sales_sheet():
    expected_headers = [
        "date",
        "customer",
        "product",
        "quantity",
        "unit_price",
        "revenue",
    ]
    report = create_excel_report(data, summary)
    workbook = load_workbook(report)
    sheet = workbook["Sales"]
    headers = list(next(sheet.iter_rows(max_row=1, values_only=True)))
    assert headers == expected_headers


def test_create_excel_report_returns_total_revenue_for_sale():
    report = create_excel_report(data, summary)
    workbook = load_workbook(report)
    sheet = workbook["Sales"]
    assert sheet["F2"].value == 360.0
    assert sheet["F3"].value == 360.0


def test_create_excel_report_creates_proper_values_in_summary_sheet():
    report = create_excel_report(data, summary)
    workbook = load_workbook(report)
    assert workbook["Summary"]["B2"].value == 4648.0
    assert workbook["Summary"]["B3"].value == 12
    assert workbook["Summary"]["B4"].value == 38
