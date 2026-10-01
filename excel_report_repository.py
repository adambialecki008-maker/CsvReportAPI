from io import BytesIO
from openpyxl import Workbook
from csv_file_repository import read_csv_file
from summary import build_summary_data, count_total_revenue_for_one_sale
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment


def create_excel_report(data, summary):
    workbook = Workbook()
    output = BytesIO()
    create_sales_sheet(workbook, data)
    create_summary_sheet(workbook, data)
    workbook.save(output)
    output.seek(0)
    return output


def create_sales_sheet(workbook, data):
    worksheet = workbook.active
    worksheet.title = "Sales"

    sales_data = data.copy()
    sales_data["revenue"] = sales_data.apply(
        lambda row: count_total_revenue_for_one_sale(
            row["quantity"],
            row["unit_price"],
        ),
        axis=1,
    )

    for row in dataframe_to_rows(
        sales_data,
        index=False,
        header=True,
    ):
        worksheet.append(row)

    style_header_rows(worksheet, [1])

    return worksheet


def create_summary_sheet(workbook, data):
    worksheet = workbook.create_sheet("Summary")
    header_rows = []
    summary = build_summary_data(data)
    # Main metrics
    header_rows.append(worksheet.max_row + 1)
    worksheet.append(["Metric", "Value"])
    worksheet.append(["Total revenue", summary["total_revenue"]])
    worksheet.append(["Orders count", summary["orders_count"]])
    worksheet.append(["Units sold", summary["units_sold"]])
    worksheet.append([])
    # Revenue per client
    header_rows.append(worksheet.max_row + 1)
    worksheet.append(["Revenue per client", "Value"])

    for client, revenue in summary["revenue_per_client"].items():
        worksheet.append([client, revenue])

    worksheet.append([])

    # Revenue per product
    header_rows.append(worksheet.max_row + 1)
    worksheet.append(["Revenue per product", "Value"])

    for product, revenue in summary["revenue_per_product"].items():
        worksheet.append([product, revenue])

    style_header_rows(worksheet, header_rows)

    return worksheet


def style_header_rows(worksheet, header_rows):
    fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    font = Font(
        bold=True,
        color="FFFFFF",
    )

    alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    for row_number in header_rows:
        for cell in worksheet[row_number]:
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment
